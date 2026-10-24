import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as L

d = pd.read_csv('data/processed/orders_clean.csv', dtype={'Postal Code': str}, parse_dates=['Order Date', 'Ship Date'])
ret = pd.read_csv('data/raw/returns.csv')[['Order ID']].drop_duplicates()
ppl = pd.read_csv('data/raw/people.csv')
N = len(d); LAST = N + 1
py_tot = pd.read_csv('data/processed/python_summary_totals.csv').set_index('Metric')['Value']

FONT = 'Arial'
NAVY, BLUE, ORANGE, LGREY, MGREY = '1F4E79', '2E75B6', 'E07B39', 'F2F2F2', 'BFBFBF'
f  = lambda **k: Font(name=FONT, size=k.pop('size', 10), **k)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color=MGREY)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
USD, USD2, PCT, INT = '$#,##0;[Red]-$#,##0', '$#,##0.00', '0.0%;[Red]-0.0%', '#,##0'

wb = Workbook()
ws_readme = wb.active; ws_readme.title = 'README'
ws_dash = wb.create_sheet('Dashboard'); ws_an = wb.create_sheet('Analysis'); ws_ins = wb.create_sheet('Insights')
ws_data = wb.create_sheet('Data'); ws_ret = wb.create_sheet('Returns'); ws_ppl = wb.create_sheet('People'); ws_chk = wb.create_sheet('Checks')

def hdr(ws, row, col, labels, color=NAVY):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(bold=True, color='FFFFFF'); c.fill = fill(color); c.border = box
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# ---------------- Returns / People lookup sheets ----------------
hdr(ws_ret, 1, 1, ['Order ID', 'Returned'])
for i, oid in enumerate(ret['Order ID'], start=2):
    ws_ret.cell(row=i, column=1, value=oid).font = f(); ws_ret.cell(row=i, column=2, value='Yes').font = f()
RET_LAST = len(ret) + 1
hdr(ws_ppl, 1, 1, ['Regional Manager', 'Region'])
for i, r in enumerate(ppl.itertuples(index=False), start=2):
    ws_ppl.cell(row=i, column=1, value=r.Person).font = f(); ws_ppl.cell(row=i, column=2, value=r.Region).font = f()
for ws, w in ((ws_ret, (18, 12)), (ws_ppl, (22, 12))):
    for i, x in enumerate(w, 1): ws.column_dimensions[L(i)].width = x

# ---------------- Data sheet ----------------
src_cols = ['Row ID', 'Order ID', 'Order Date', 'Ship Date', 'Ship Mode', 'Customer ID', 'Customer Name', 'Segment', 'State', 'City',
            'Postal Code', 'Region', 'Product ID', 'Category', 'Sub-Category', 'Product Name', 'Sales', 'Quantity', 'Discount', 'Profit']
calc_cols = ['Order Year', 'Order Month', 'Profit Margin', 'Discount Band', 'Ship Days', 'Returned', 'Regional Manager', 'New Order Flag']
hdr(ws_data, 1, 1, src_cols); hdr(ws_data, 1, len(src_cols) + 1, calc_cols, color=ORANGE)
ws_data.row_dimensions[1].height = 30
for r, row in enumerate(d[src_cols].itertuples(index=False), start=2):
    for c, v in enumerate(row, start=1):
        if c in (3, 4): v = v.to_pydatetime()
        elif c in (17, 18, 19, 20): v = float(v) if c != 18 else int(v)
        elif c == 1: v = int(v)
        cell = ws_data.cell(row=r, column=c, value=v)
    # calculated columns (formulas)
    ws_data.cell(row=r, column=21, value=f'=YEAR(C{r})')
    ws_data.cell(row=r, column=22, value=f'=MONTH(C{r})')
    ws_data.cell(row=r, column=23, value=f'=IF(Q{r}=0,0,T{r}/Q{r})')
    ws_data.cell(row=r, column=24, value=f'=IF(S{r}=0,"0%",IF(S{r}<=0.1,"0-10%",IF(S{r}<=0.2,"10-20%",IF(S{r}<=0.3,"20-30%",IF(S{r}<=0.5,"30-50%","50%+")))))')
    ws_data.cell(row=r, column=25, value=f'=D{r}-C{r}')
    ws_data.cell(row=r, column=26, value=f'=IF(ISNA(MATCH(B{r},Returns!$A$2:$A${RET_LAST},0)),"No","Yes")')
    ws_data.cell(row=r, column=27, value=f'=INDEX(People!$A$2:$A$5,MATCH(L{r},People!$B$2:$B$5,0))')
    ws_data.cell(row=r, column=28, value=f'=IF(B{r}=B{r-1},0,1)')
fmt = {3: 'yyyy-mm-dd', 4: 'yyyy-mm-dd', 17: '#,##0.00', 19: '0%', 20: '#,##0.00', 23: '0.0%', 11: '@'}
for col, nf in fmt.items():
    for r in range(2, LAST + 1): ws_data.cell(row=r, column=col).number_format = nf
widths = [8, 16, 12, 12, 15, 12, 20, 13, 16, 16, 11, 9, 17, 16, 13, 45, 11, 9, 9, 11, 10, 10, 11, 12, 10, 10, 18, 11]
for i, w in enumerate(widths, 1): ws_data.column_dimensions[L(i)].width = w
for row in ws_data.iter_rows(min_row=2, max_row=LAST, max_col=28):
    for c in row: c.font = f()
ws_data.freeze_panes = 'C2'; ws_data.auto_filter.ref = f'A1:{L(28)}{LAST}'

# ---------------- Named ranges ----------------
names = {'d_OrderID': 'B', 'd_Segment': 'H', 'd_Region': 'L', 'd_Category': 'N', 'd_SubCat': 'O', 'd_Product': 'P', 'd_Sales': 'Q', 'd_Qty': 'R',
         'd_Discount': 'S', 'd_Profit': 'T', 'd_Year': 'U', 'd_Month': 'V', 'd_Band': 'X', 'd_Returned': 'Z', 'd_NewOrder': 'AB'}
for n, col in names.items():
    wb.defined_names[n] = DefinedName(n, attr_text=f'Data!${col}$2:${col}${LAST}')
wb.defined_names['f_Year'] = DefinedName('f_Year', attr_text='Analysis!$B$3')
wb.defined_names['f_Region'] = DefinedName('f_Region', attr_text='Analysis!$B$4')

# ---------------- Analysis sheet ----------------
A = ws_an
A['A1'] = 'Analysis tables (formula-driven from the Data sheet)'; A['A1'].font = f(bold=True, size=14, color=NAVY)
A['A3'] = 'Year filter criterion'; A['B3'] = '=IF(Dashboard!$C$5="All",">0",Dashboard!$C$5)'
A['A4'] = 'Region filter criterion'; A['B4'] = '=IF(Dashboard!$F$5="All","*",Dashboard!$F$5)'
A['C3'] = 'Set on the Dashboard sheet (yellow cells). ">0" and "*" mean "all".'; A['C3'].font = f(italic=True, color='7F7F7F')
for c in ('A3', 'A4'): A[c].font = f(bold=True)
for c in ('B3', 'B4'): A[c].font = f(color='008000'); A[c].alignment = Alignment(horizontal='center')

YR, RG = 'd_Year,f_Year', 'd_Region,f_Region'
pos = {}
def title(row, text, note):
    A.cell(row=row, column=1, value=text).font = f(bold=True, size=12, color=NAVY)
    A.cell(row=row, column=4 if False else 8, value=note).font = f(italic=True, color='7F7F7F')

# 1 Yearly (no filters)
r = 6; title(r, '1. Yearly summary', 'No filters (all regions)')
hdr(A, r + 1, 1, ['Year', 'Sales', 'Profit', 'Profit Margin', 'Orders', 'Sales YoY %'])
for i, y in enumerate([2015, 2016, 2017, 2018]):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=y)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Year,$A{rr})')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_Year,$A{rr})')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
    A.cell(row=rr, column=5, value=f'=SUMIFS(d_NewOrder,d_Year,$A{rr})')
    A.cell(row=rr, column=6, value='' if i == 0 else f'=B{rr}/B{rr-1}-1')
pos['yearly'] = (r + 2, r + 5)
rr = r + 6
A.cell(row=rr, column=1, value='Total'); A.cell(row=rr, column=2, value=f'=SUM(B{r+2}:B{r+5})'); A.cell(row=rr, column=3, value=f'=SUM(C{r+2}:C{r+5})')
A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})'); A.cell(row=rr, column=5, value=f'=SUM(E{r+2}:E{r+5})')
for c in range(1, 7): A.cell(row=rr, column=c).font = f(bold=True)
pos['yearly_total'] = rr

# 2 Monthly (region filter)
r = 15; title(r, '2. Monthly trend', 'Filter: Region only (all years shown)')
hdr(A, r + 1, 1, ['Month', 'Sales', 'Profit', 'MoM Sales %', 'Year', 'Month No.'])
row = r + 2
for y in (2015, 2016, 2017, 2018):
    for m in range(1, 13):
        A.cell(row=row, column=1, value=f'{y}-{m:02d}')
        A.cell(row=row, column=5, value=y).font = f(color='7F7F7F'); A.cell(row=row, column=6, value=m).font = f(color='7F7F7F')
        A.cell(row=row, column=2, value=f'=SUMIFS(d_Sales,d_Year,$E{row},d_Month,$F{row},d_Region,f_Region)')
        A.cell(row=row, column=3, value=f'=SUMIFS(d_Profit,d_Year,$E{row},d_Month,$F{row},d_Region,f_Region)')
        A.cell(row=row, column=4, value='' if row == r + 2 else f'=IF(B{row-1}=0,"",B{row}/B{row-1}-1)')
        row += 1
pos['monthly'] = (r + 2, row - 1)
mon_first, mon_last = pos['monthly']

# 3 Calendar-month seasonality (year + region)
r = row + 2; title(r, '3. Seasonality: sales by calendar month', 'Filters: Year + Region')
hdr(A, r + 1, 1, ['Month', 'Sales', 'Share of Sales'])
mnames = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
for m in range(12):
    rr = r + 2 + m
    A.cell(row=rr, column=1, value=mnames[m])
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Month,{m+1},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=IF(SUM($B${r+2}:$B${r+13})=0,0,B{rr}/SUM($B${r+2}:$B${r+13}))')
pos['season'] = (r + 2, r + 13)

# 4 Category
r = r + 16; title(r, '4. Category performance', 'Filters: Year + Region')
cats = ['Technology', 'Office Supplies', 'Furniture']
hdr(A, r + 1, 1, ['Category', 'Sales', 'Profit', 'Profit Margin', 'Share of Sales', 'Share of Profit', 'Avg Discount'])
for i, c_ in enumerate(cats):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=c_)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Category,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_Category,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
    A.cell(row=rr, column=5, value=f'=IF(SUM($B${r+2}:$B${r+4})=0,0,B{rr}/SUM($B${r+2}:$B${r+4}))')
    A.cell(row=rr, column=6, value=f'=IF(SUM($C${r+2}:$C${r+4})=0,0,C{rr}/SUM($C${r+2}:$C${r+4}))')
    A.cell(row=rr, column=7, value=f'=IFERROR(AVERAGEIFS(d_Discount,d_Category,$A{rr},d_Year,f_Year,d_Region,f_Region),0)')
pos['cat'] = (r + 2, r + 4)

# 5 Sub-category
r = r + 7; title(r, '5. Sub-category performance', 'Filters: Year + Region (order = profit rank on full data)')
subs = d.groupby('Sub-Category').Profit.sum().sort_values(ascending=False).index.tolist()
hdr(A, r + 1, 1, ['Sub-Category', 'Sales', 'Profit', 'Profit Margin', 'Avg Discount'])
for i, s in enumerate(subs):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=s)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_SubCat,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_SubCat,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
    A.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS(d_Discount,d_SubCat,$A{rr},d_Year,f_Year,d_Region,f_Region),0)')
pos['sub'] = (r + 2, r + 1 + len(subs))

# 6 Region (year filter)
r = r + 3 + len(subs); title(r, '6. Region performance', 'Filter: Year only')
regs = ['West', 'East', 'South', 'Central']
hdr(A, r + 1, 1, ['Region', 'Sales', 'Profit', 'Profit Margin', 'Avg Discount', 'Regional Manager'])
for i, g in enumerate(regs):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=g)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Region,$A{rr},d_Year,f_Year)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_Region,$A{rr},d_Year,f_Year)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
    A.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS(d_Discount,d_Region,$A{rr},d_Year,f_Year),0)')
    A.cell(row=rr, column=6, value=f'=INDEX(People!$A$2:$A$5,MATCH(A{rr},People!$B$2:$B$5,0))')
pos['reg'] = (r + 2, r + 5)

# 7 Segment
r = r + 8; title(r, '7. Customer segment performance', 'Filters: Year + Region')
hdr(A, r + 1, 1, ['Segment', 'Sales', 'Profit', 'Profit Margin'])
for i, s in enumerate(['Consumer', 'Corporate', 'Home Office']):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=s)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Segment,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_Segment,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
pos['seg'] = (r + 2, r + 4)

# 8 Discount band
r = r + 7; title(r, '8. Discount band vs profitability', 'Filters: Year + Region')
bands = ['0%', '0-10%', '10-20%', '20-30%', '30-50%', '50%+']
hdr(A, r + 1, 1, ['Discount Band', 'Order Lines', 'Sales', 'Profit', 'Profit Margin'])
for i, b in enumerate(bands):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=b)
    A.cell(row=rr, column=2, value=f'=COUNTIFS(d_Band,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Sales,d_Band,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=SUMIFS(d_Profit,d_Band,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=5, value=f'=IF(C{rr}=0,0,D{rr}/C{rr})')
pos['band'] = (r + 2, r + 7)
A.cell(row=r + 8, column=1, value='Bands: 0% = no discount; 0-10% means above 0% up to 10%, and so on.').font = f(italic=True, color='7F7F7F')

# 9 Top 10 products
r = r + 11; title(r, '9. Top 10 products by sales (full data)', 'Names fixed from full data; values respond to Year + Region')
top = d.groupby('Product Name').Sales.sum().nlargest(10).index.tolist()
hdr(A, r + 1, 1, ['Product Name', 'Sales', 'Profit', 'Profit Margin', 'Chart Label'])
for i, p in enumerate(top):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=p)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Product,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Profit,d_Product,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
    A.cell(row=rr, column=5, value=(p if len(p) <= 28 else p[:27] + '...'))
pos['top'] = (r + 2, r + 11)

# 10 Returns
r = r + 14; title(r, '10. Returns', 'Filters: Year + Region')
hdr(A, r + 1, 1, ['Category', 'Sales', 'Returned Sales', 'Returned Sales %'])
for i, c_ in enumerate(cats):
    rr = r + 2 + i
    A.cell(row=rr, column=1, value=c_)
    A.cell(row=rr, column=2, value=f'=SUMIFS(d_Sales,d_Category,$A{rr},d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=3, value=f'=SUMIFS(d_Sales,d_Category,$A{rr},d_Returned,"Yes",d_Year,f_Year,d_Region,f_Region)')
    A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
rr = r + 5
A.cell(row=rr, column=1, value='All categories'); A.cell(row=rr, column=2, value=f'=SUM(B{r+2}:B{r+4})'); A.cell(row=rr, column=3, value=f'=SUM(C{r+2}:C{r+4})')
A.cell(row=rr, column=4, value=f'=IF(B{rr}=0,0,C{rr}/B{rr})')
A.cell(row=rr + 1, column=1, value='Returned orders % of orders')
A.cell(row=rr + 1, column=4, value=f'=IFERROR(SUMIFS(d_NewOrder,d_Returned,"Yes",d_Year,f_Year,d_Region,f_Region)/SUMIFS(d_NewOrder,d_Year,f_Year,d_Region,f_Region),0)')
for c in range(1, 5): A.cell(row=rr, column=c).font = f(bold=True)
pos['ret'] = (r + 2, rr + 1)

# formats for Analysis
money_cols = {'yearly': (2, 3), 'monthly': (2, 3), 'season': (2,), 'cat': (2, 3), 'sub': (2, 3), 'reg': (2, 3), 'seg': (2, 3), 'band': (3, 4), 'top': (2, 3), 'ret': (2, 3)}
pct_cols = {'yearly': (4, 6), 'monthly': (4,), 'season': (3,), 'cat': (4, 5, 6, 7), 'sub': (4, 5), 'reg': (4, 5), 'seg': (4,), 'band': (5,), 'top': (4,), 'ret': (4,)}
for k, v_ in pos.items():
    if k == 'yearly_total': continue
    a, b = v_
    for rr in range(a, b + 1):
        for c in money_cols.get(k, ()): A.cell(row=rr, column=c).number_format = USD
        for c in pct_cols.get(k, ()): A.cell(row=rr, column=c).number_format = PCT
for rr in range(pos['yearly'][0], pos['yearly_total'] + 1):
    A.cell(row=rr, column=2).number_format = USD; A.cell(row=rr, column=3).number_format = USD; A.cell(row=rr, column=5).number_format = INT
    A.cell(row=rr, column=4).number_format = PCT; A.cell(row=rr, column=6).number_format = PCT
for rr in range(pos['band'][0], pos['band'][1] + 1): A.cell(row=rr, column=2).number_format = INT
a, b = pos['ret']; A.cell(row=b, column=4).number_format = PCT
for col, w in zip('ABCDEFGH', (42, 14, 14, 14, 14, 20, 14, 40)): A.column_dimensions[col].width = w
A.column_dimensions['E'].width = 30
for key in ('sub', 'cat', 'reg', 'seg', 'band', 'top'):
    a_, b_ = pos[key]
    col = {'sub': 'D', 'cat': 'D', 'reg': 'D', 'seg': 'D', 'band': 'E', 'top': 'D'}[key]
    A.conditional_formatting.add(f'{col}{a_}:{col}{b_}', CellIsRule(operator='lessThan', formula=['0'], fill=fill('FADBD8')))
A.freeze_panes = 'A5'

# ---------------- Dashboard ----------------
D = ws_dash; D.sheet_view.showGridLines = False
D.column_dimensions['A'].width = 2; D.column_dimensions['N'].width = 2
for c in range(2, 14): D.column_dimensions[L(c)].width = 12.5
D.merge_cells('B2:M2'); D['B2'] = 'Retail Sales Performance Dashboard'; D['B2'].font = f(bold=True, size=20, color=NAVY)
D.merge_cells('B3:M3'); D['B3'] = 'Sample Superstore data, US retail orders 2015-2018. Pick a Year and Region in the yellow cells; KPIs and charts update.'
D['B3'].font = f(italic=True, color='595959')
D['B5'] = 'Year'; D['E5'] = 'Region'
for c in ('B5', 'E5'): D[c].font = f(bold=True); D[c].alignment = Alignment(horizontal='right')
for c, v in (('C5', 'All'), ('F5', 'All')):
    D[c] = v; D[c].fill = fill('FFFF00'); D[c].font = f(bold=True, color='0000FF'); D[c].alignment = Alignment(horizontal='center'); D[c].border = box
dv1 = DataValidation(type='list', formula1='"All,2015,2016,2017,2018"', allow_blank=False); dv2 = DataValidation(type='list', formula1='"All,Central,East,South,West"', allow_blank=False)
D.add_data_validation(dv1); D.add_data_validation(dv2); dv1.add('C5'); dv2.add('F5')
D.merge_cells('H5:M5'); D['H5'] = 'Monthly chart: Region filter only. Region chart: Year filter only.'; D['H5'].font = f(italic=True, color='7F7F7F', size=9)
kpis = [('Total Sales', '=SUMIFS(d_Sales,d_Year,f_Year,d_Region,f_Region)', USD),
        ('Total Profit', '=SUMIFS(d_Profit,d_Year,f_Year,d_Region,f_Region)', USD),
        ('Profit Margin', '=IF(B8=0,0,D8/B8)', PCT),
        ('Orders', '=SUMIFS(d_NewOrder,d_Year,f_Year,d_Region,f_Region)', INT),
        ('Avg Order Value', '=IF(J8=0,0,B8/J8)', USD2),
        ('Units Sold', '=SUMIFS(d_Qty,d_Year,f_Year,d_Region,f_Region)', INT)]
for i, (lab, frm, nf) in enumerate(kpis):
    c1, c2 = 2 + 2 * i, 3 + 2 * i
    D.merge_cells(start_row=7, start_column=c1, end_row=7, end_column=c2); D.merge_cells(start_row=8, start_column=c1, end_row=8, end_column=c2)
    a_ = D.cell(row=7, column=c1, value=lab); a_.font = f(bold=True, size=9, color='FFFFFF'); a_.fill = fill(NAVY); a_.alignment = Alignment(horizontal='center')
    v_ = D.cell(row=8, column=c1, value=frm); v_.font = f(bold=True, size=16, color=NAVY); v_.number_format = nf; v_.alignment = Alignment(horizontal='center', vertical='center')
    for rr in (7, 8):
        for cc in (c1, c2): D.cell(row=rr, column=cc).fill = fill(NAVY if rr == 7 else LGREY); D.cell(row=rr, column=cc).border = box
D.row_dimensions[8].height = 34
# KPI formulas reference cells by position: fix to real positions (Sales B8, Profit D8, Orders J8)
D['F8'] = '=IF(B8=0,0,D8/B8)'; D['L8'] = '=SUMIFS(d_Qty,d_Year,f_Year,d_Region,f_Region)'; D['H8'] = '=IF(B8=0,0,SUMIFS(d_Sales,d_Year,f_Year,d_Region,f_Region)/SUMIFS(d_NewOrder,d_Year,f_Year,d_Region,f_Region))'
D['J8'] = '=SUMIFS(d_NewOrder,d_Year,f_Year,d_Region,f_Region)'
# order of cards: Sales(B) Profit(D) Margin(F) AOV(H) Orders(J) Units(L)
for col, (lab, nf) in zip((2, 4, 6, 8, 10, 12), (('Total Sales', USD), ('Total Profit', USD), ('Profit Margin', PCT), ('Avg Order Value', USD2), ('Orders', INT), ('Units Sold', INT))):
    D.cell(row=7, column=col, value=lab); D.cell(row=8, column=col).number_format = nf

def style(ch, title_, w, h):
    ch.title = title_; ch.width = w; ch.height = h
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.legend.position = 'b'
def color(series, hexc, line=False):
    if line: series.graphicalProperties.line.solidFill = hexc; series.graphicalProperties.line.width = 22000; series.smooth = False
    else: series.graphicalProperties.solidFill = hexc; series.graphicalProperties.line.solidFill = hexc

# monthly line
a_, b_ = pos['monthly']
ch = LineChart(); style(ch, 'Monthly Sales and Profit', 29, 7.5)
ch.add_data(Reference(A, min_col=2, max_col=3, min_row=a_ - 1, max_row=b_), titles_from_data=True)
ch.set_categories(Reference(A, min_col=1, min_row=a_, max_row=b_))
color(ch.series[0], BLUE, True); color(ch.series[1], ORANGE, True)
ch.y_axis.number_format = '$#,##0'; ch.x_axis.tickLblSkip = 3; ch.x_axis.tickMarkSkip = 3; ch.x_axis.tickLblPos = 'low'
D.add_chart(ch, 'B10')
# category
a_, b_ = pos['cat']
ch = BarChart(); ch.type = 'col'; style(ch, 'Sales and Profit by Category', 14.3, 7.5)
ch.add_data(Reference(A, min_col=2, max_col=3, min_row=a_ - 1, max_row=b_), titles_from_data=True); ch.set_categories(Reference(A, min_col=1, min_row=a_, max_row=b_))
color(ch.series[0], BLUE); color(ch.series[1], ORANGE); ch.y_axis.number_format = '$#,##0'
D.add_chart(ch, 'B25')
# region
a_, b_ = pos['reg']
ch = BarChart(); ch.type = 'col'; style(ch, 'Sales and Profit by Region', 14.3, 7.5)
ch.add_data(Reference(A, min_col=2, max_col=3, min_row=a_ - 1, max_row=b_), titles_from_data=True); ch.set_categories(Reference(A, min_col=1, min_row=a_, max_row=b_))
color(ch.series[0], BLUE); color(ch.series[1], ORANGE); ch.y_axis.number_format = '$#,##0'
D.add_chart(ch, 'H25')
# discount band margin
a_, b_ = pos['band']
ch = BarChart(); ch.type = 'col'; style(ch, 'Profit Margin by Discount Band', 14.3, 8.5)
ch.add_data(Reference(A, min_col=5, min_row=a_ - 1, max_row=b_), titles_from_data=True); ch.set_categories(Reference(A, min_col=1, min_row=a_, max_row=b_))
color(ch.series[0], ORANGE); ch.y_axis.number_format = '0%'; ch.legend = None
ch.x_axis.tickLblPos = 'low'
ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True; ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False; ch.dataLabels.showPercent = False; ch.dataLabels.numFmt = '0%'
D.add_chart(ch, 'B40')
# top 10 products
a_, b_ = pos['top']
ch = BarChart(); ch.type = 'bar'; style(ch, 'Top 10 Products by Sales', 14.3, 8.5)
ch.add_data(Reference(A, min_col=2, max_col=3, min_row=a_ - 1, max_row=b_), titles_from_data=True); ch.set_categories(Reference(A, min_col=5, min_row=a_, max_row=b_))
color(ch.series[0], BLUE); color(ch.series[1], ORANGE); ch.x_axis.scaling.orientation = 'maxMin'; ch.x_axis.tickLblPos = 'low'; ch.y_axis.number_format = '$#,##0'
D.add_chart(ch, 'H40')
D.page_setup.orientation = 'portrait'; D.page_setup.fitToWidth = 1; D.page_setup.fitToHeight = 1; D.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------- Insights ----------------
I = ws_ins; I.sheet_view.showGridLines = False
I['A1'] = 'Business insights and recommendations'; I['A1'].font = f(bold=True, size=14, color=NAVY)
I['A2'] = 'Figures are for the full dataset (Year = All, Region = All). Use the Dashboard filters to test each point. These are associations in one sample dataset, not proof of cause.'
I['A2'].font = f(italic=True, color='595959')
hdr(I, 4, 1, ['#', 'Finding', 'Evidence', 'Recommendation'])
ins = [
 ('Sales and margin are growing.', 'Sales rose from $484K (2015) to $733K (2018), +29% in 2017 and +20% in 2018. Profit margin moved from 10.2% to 12.7%.', 'Keep the 2017-2018 growth drivers. Check which categories and regions produced the growth before setting 2019 targets.'),
 ('Sales are concentrated in three months.', 'September, November and December bring about 43% of sales. January and February are the weakest months.', 'Build stock and staffing plans around Sep-Dec. Test promotions in Jan-Feb to smooth demand.'),
 ('Discounts above 20% lose money.', 'Every discount band up to 20% is profitable (margin 30% at no discount, 12% at 10-20%). Bands above 20% have negative margins, down to -119% at 50%+. Lines discounted above 20% (14% of lines) lost about $135K, against $286K total profit.', 'Cap standard discounts at 20% and require approval above that. Review the products and customers that receive the deepest discounts.'),
 ('Furniture sells a lot but earns little.', 'Furniture is 32% of sales but 6% of profit (margin 2.5%). Tables lose about $18K and Bookcases about $3K, with average discounts of 26% and 21%.', 'Review pricing and discount rules for Tables and Bookcases first. Shift sales effort toward Technology and Office Supplies, which earn a 17% margin.'),
 ('The Central region has the weakest margin.', 'Central margin is 7.9% against 14.9% in the West. Central also has the highest average discount (24% against 11% in the West).', 'Look at discount approval in Central. Compare its product mix with the West to see how much comes from discounting versus mix.'),
 ('Returns take about 8% of sales value.', 'About 5.9% of orders and 7.9% of sales value are returned. Returned sales are highest in Technology (8.7%).', 'Find the reasons for returns on high-value Technology orders. Even a small reduction protects profit.'),
]
for i, (a, b, c) in enumerate(ins, start=1):
    for col, v in enumerate((i, a, b, c), start=1):
        cell = I.cell(row=4 + i, column=col, value=v); cell.font = f(bold=(col == 2)); cell.alignment = Alignment(wrap_text=True, vertical='top'); cell.border = box
    I.row_dimensions[4 + i].height = 78
for col, w in zip('ABCD', (4, 30, 70, 55)): I.column_dimensions[col].width = w

# ---------------- Checks ----------------
K = ws_chk
K['A1'] = 'Reconciliation: Excel formulas vs Python (pandas) results'; K['A1'].font = f(bold=True, size=14, color=NAVY)
K['A2'] = 'Blue values are typed in from the Python notebook output (data/processed/python_summary_totals.csv). Black values are live Excel formulas on the Data sheet.'
K['A2'].font = f(italic=True, color='595959')
hdr(K, 4, 1, ['Check', 'Excel', 'Python', 'Difference', 'Status'])
rows = [('Row count after cleaning', f'=COUNTA(Data!B2:B{LAST})', py_tot['Rows'], INT),
        ('Distinct orders', '=SUM(d_NewOrder)', py_tot['Orders'], INT),
        ('Total sales', '=SUM(d_Sales)', py_tot['Total Sales'], USD2),
        ('Total profit', '=SUM(d_Profit)', py_tot['Total Profit'], USD2),
        ('Profit margin %', '=SUM(d_Profit)/SUM(d_Sales)*100', py_tot['Profit Margin %'], '0.00'),
        ('Yearly table total sales = Data total sales', f'=Analysis!B{pos["yearly_total"]}', None, USD2)]
for i, (lab, frm, pyv, nf) in enumerate(rows, start=5):
    K.cell(row=i, column=1, value=lab); K.cell(row=i, column=2, value=frm)
    if pyv is None: K.cell(row=i, column=3, value='=SUM(d_Sales)')
    else: K.cell(row=i, column=3, value=float(pyv))
    K.cell(row=i, column=4, value=f'=B{i}-C{i}'); K.cell(row=i, column=5, value=f'=IF(ABS(D{i})<0.01,"OK","CHECK")')
    for c in range(1, 6): K.cell(row=i, column=c).font = f(color='0000FF' if (c == 3 and pyv is not None) else '000000'); K.cell(row=i, column=c).border = box
    for c in (2, 3, 4): K.cell(row=i, column=c).number_format = nf
K['A12'] = 'Note: Python values were rounded to 2 decimals when saved, so differences below 0.01 are treated as a match.'; K['A12'].font = f(italic=True, color='7F7F7F')
K['A13'] = 'Assumption: the Data sheet keeps the original order of lines, so all lines of one order are next to each other (verified: 5,009 blocks = 5,009 distinct orders). The New Order Flag column depends on this.'
K['A13'].font = f(italic=True, color='7F7F7F')
for col, w in zip('ABCDE', (46, 18, 18, 14, 10)): K.column_dimensions[col].width = w

# ---------------- README ----------------
R = ws_readme; R.sheet_view.showGridLines = False
R['A1'] = 'Retail Sales Performance Analysis: Excel workbook'; R['A1'].font = f(bold=True, size=16, color=NAVY)
lines = [
 ('Purpose', 'Clean and analyze a public retail dataset, and report revenue, profit, discounts, regions and returns with formulas, charts and written insights.'),
 ('Data source', 'Sample Superstore (Tableau sample data): Orders, People and Returns tables. US orders, Jan 2015 to Dec 2018. Public sample data, not a real company.'),
 ('', ''),
 ('SHEETS', ''),
 ('Dashboard', 'KPIs and charts. Change the yellow Year and Region cells to filter.'),
 ('Analysis', 'Ten formula-driven summary tables (SUMIFS, COUNTIFS, AVERAGEIFS) that feed the Dashboard.'),
 ('Insights', 'Six findings with evidence and a recommendation each.'),
 ('Data', 'Cleaned order lines (9,993 rows). Columns A-T are source data; the orange columns U-AB are calculated with formulas.'),
 ('Returns / People', 'Lookup tables used by INDEX/MATCH formulas in the Data sheet.'),
 ('Checks', 'Reconciles Excel totals to the Python notebook results.'),
 ('', ''),
 ('CLEANING DONE (in Python, see notebook)', ''),
 ('1', 'Converted dates from text to real dates.'),
 ('2', 'Removed 1 exact duplicate order line (Order US-2015-150119, sales $281.37). Raw 9,994 rows became 9,993.'),
 ('3', 'Filled 11 missing Postal Codes (all Burlington, Vermont) with 05401. Postal codes are kept as text.'),
 ('4', 'Returns table: 800 rows with repeated order IDs reduced to 296 unique returned orders.'),
 ('5', 'Dropped the Country column (always United States).'),
 ('', ''),
 ('CALCULATED COLUMNS (Data sheet)', ''),
 ('Order Year / Order Month', 'YEAR() and MONTH() of Order Date.'),
 ('Profit Margin', 'Profit / Sales.'),
 ('Discount Band', 'Nested IF: 0%, 0-10%, 10-20%, 20-30%, 30-50%, 50%+ (upper edge included in each band).'),
 ('Ship Days', 'Ship Date minus Order Date.'),
 ('Returned', 'INDEX/MATCH style lookup of Order ID in the Returns sheet (Yes/No).'),
 ('Regional Manager', 'INDEX/MATCH of Region in the People sheet.'),
 ('New Order Flag', '1 on the first line of each order, 0 on the rest. Summing it gives the distinct order count.'),
 ('', ''),
 ('NOTES', ''),
 ('Colors', 'Yellow cells are inputs you can change. Blue text on the Checks sheet marks values typed in from Python.'),
 ('Not included', 'This file uses formula-based summary tables. Real PivotTables and slicers are built from the Data sheet by following docs/excel_pivot_guide.md.'),
]
for i, (a, b) in enumerate(lines, start=3):
    ca, cb = R.cell(row=i, column=1, value=a), R.cell(row=i, column=2, value=b)
    ca.font = f(bold=True, color=NAVY if (a.isupper() or a == 'SHEETS') else '000000'); cb.font = f()
    cb.alignment = Alignment(wrap_text=True, vertical='top'); ca.alignment = Alignment(vertical='top', wrap_text=True)
R.column_dimensions['A'].width = 34; R.column_dimensions['B'].width = 110
R.page_setup.orientation = 'landscape'; R.page_setup.fitToWidth = 1; R.page_setup.fitToHeight = 0; R.sheet_properties.pageSetUpPr = __import__('openpyxl').worksheet.properties.PageSetupProperties(fitToPage=True)

# default fonts for any remaining un-styled cells in Analysis
for row in A.iter_rows(min_row=1, max_row=A.max_row, max_col=8):
    for c in row:
        if c.value is not None and (c.font is None or c.font.name != FONT): c.font = f()
for ws in (ws_ret, ws_ppl):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=2):
        for c in row:
            if c.font.name != FONT: c.font = f()
wb.calculation.fullCalcOnLoad = True
wb.save('excel/Retail_Sales_Analysis.xlsx')
print('saved', pos)
