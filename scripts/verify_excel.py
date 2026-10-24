"""Check excel/Retail_Sales_Analysis.xlsx against data/processed/orders_clean.csv.

Usage (run from the project root, with python, not as a shell script):
    python scripts/verify_excel.py default
        Checks the workbook as saved, with Year = All and Region = All.

    python scripts/verify_excel.py manual <xlsx file> <year> <region>
        Checks a copy of the workbook saved from Excel after you changed the
        Year (Dashboard!C5) and Region (Dashboard!F5) filters. Example:
        python scripts/verify_excel.py manual excel/test_filtered.xlsx 2017 West
        Excel recalculates when you save, so the file holds the new values.
"""
import sys
import pandas as pd
from openpyxl import load_workbook

WORKBOOK = 'excel/Retail_Sales_Analysis.xlsx'

d = pd.read_csv('data/processed/orders_clean.csv', parse_dates=['Order Date'])
d['Y'] = d['Order Date'].dt.year
d['M'] = d['Order Date'].dt.month

NCHK = 0


def expected(year, region):
    x = d.copy()
    if year != 'All':
        x = x[x.Y == int(year)]
    if region != 'All':
        x = x[x.Region == region]
    return x


def run(path, year, region):
    global NCHK
    NCHK = 0
    wb = load_workbook(path, data_only=True)
    D, A = wb['Dashboard'], wb['Analysis']
    xy = expected(year, region)          # both filters
    yr_only = expected(year, 'All')      # region table ignores the region filter
    rg_only = expected('All', region)    # month table ignores the year filter
    bad = []

    def chk(name, got, exp, tol=0.01):
        global NCHK
        NCHK += 1
        if got is None or abs(float(got) - float(exp)) > tol:
            bad.append((name, got, exp))

    # KPI cards
    s, p = xy.Sales.sum(), xy.Profit.sum()
    o = xy['Order ID'].nunique()
    chk('sales', D['B8'].value, s)
    chk('profit', D['D8'].value, p)
    chk('margin', D['F8'].value, p / s, 1e-6)
    chk('AOV', D['H8'].value, s / o)
    chk('orders', D['J8'].value, o)
    chk('units', D['L8'].value, xy.Quantity.sum())

    # Tables on the Analysis sheet, found by scanning column A for labels
    col = {r: A.cell(row=r, column=1).value for r in range(1, A.max_row + 1)}
    sub_categories = set(d['Sub-Category'])
    managers = ('Anna Andreadi', 'Chuck Magee', 'Kelly Williams', 'Cassandra Brandow')

    for r in range(1, A.max_row + 1):
        v = col[r]

        # category
        if v in ('Technology', 'Office Supplies', 'Furniture') and A.cell(row=r, column=7).value is not None:
            g = xy[xy.Category == v]
            chk(f'cat {v} sales', A.cell(row=r, column=2).value, g.Sales.sum())
            chk(f'cat {v} profit', A.cell(row=r, column=3).value, g.Profit.sum())
            chk(f'cat {v} avgdisc', A.cell(row=r, column=7).value, g.Discount.mean() if len(g) else 0, 1e-6)

        # region (ignores the region filter, respects the year filter)
        if v in ('West', 'East', 'South', 'Central') and A.cell(row=r, column=6).value in managers:
            g = yr_only[yr_only.Region == v]
            chk(f'reg {v} sales', A.cell(row=r, column=2).value, g.Sales.sum())
            chk(f'reg {v} profit', A.cell(row=r, column=3).value, g.Profit.sum())

        # sub-category
        if isinstance(v, str) and v in sub_categories:
            g = xy[xy['Sub-Category'] == v]
            chk(f'sub {v} sales', A.cell(row=r, column=2).value, g.Sales.sum())
            chk(f'sub {v} profit', A.cell(row=r, column=3).value, g.Profit.sum())

        # discount bands
        if (v in ('0%', '0-10%', '10-20%', '20-30%', '30-50%', '50%+')
                and A.cell(row=r, column=5).value is not None
                and isinstance(A.cell(row=r, column=2).value, (int, float))):
            g = xy[xy['Discount Band'] == v]
            chk(f'band {v} lines', A.cell(row=r, column=2).value, len(g))
            chk(f'band {v} profit', A.cell(row=r, column=4).value, g.Profit.sum())

        # segment
        if v in ('Consumer', 'Corporate', 'Home Office'):
            g = xy[xy.Segment == v]
            chk(f'seg {v} sales', A.cell(row=r, column=2).value, g.Sales.sum())

        # month (labels like 2017-03; ignores the year filter, respects the region filter)
        if isinstance(v, str) and len(v) == 7 and v[4] == '-' and v[:4].isdigit():
            g = rg_only[(rg_only.Y == int(v[:4])) & (rg_only.M == int(v[5:]))]
            chk(f'month {v} sales', A.cell(row=r, column=2).value, g.Sales.sum())
            chk(f'month {v} profit', A.cell(row=r, column=3).value, g.Profit.sum())

    # returns
    for r in range(1, A.max_row + 1):
        if col[r] == 'All categories':
            returned = xy[xy.Returned == 'Yes']
            chk('returned sales', A.cell(row=r, column=3).value, returned.Sales.sum())
            chk('returned order %', A.cell(row=r + 1, column=4).value,
                returned['Order ID'].nunique() / xy['Order ID'].nunique() if len(xy) else 0, 1e-6)

    # top 10 products by overall sales
    top = d.groupby('Product Name').Sales.sum().nlargest(10).index.tolist()
    for r in range(1, A.max_row + 1):
        if col[r] in top:
            g = xy[xy['Product Name'] == col[r]]
            chk('top sales ' + col[r][:20], A.cell(row=r, column=2).value, g.Sales.sum())

    # status cells on the Checks sheet
    K = wb['Checks']
    stat = [K.cell(row=r, column=5).value for r in range(5, 11)]
    return bad, stat


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ('default', 'manual'):
        print(__doc__)
        sys.exit(1)

    mode = sys.argv[1]
    if mode == 'default':
        bad, stat = run(WORKBOOK, 'All', 'All')
        label = 'DEFAULT'
    else:
        if len(sys.argv) != 5:
            print('manual mode needs: <xlsx file> <year> <region>')
            print(__doc__)
            sys.exit(1)
        bad, stat = run(sys.argv[2], sys.argv[3], sys.argv[4])
        label = f'MANUAL ({sys.argv[3]}, {sys.argv[4]})'

    print(f'{label} mismatches:', bad, '| comparisons:', NCHK)
    print('Checks sheet:', stat)
    sys.exit(1 if bad else 0)


if __name__ == '__main__':
    main()